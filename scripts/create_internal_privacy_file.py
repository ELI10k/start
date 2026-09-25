from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

OUT = "outputs/START-Life-Fit-internal-privacy-file-he.docx"

def rtl(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = p._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1")

def para(doc, text="", bold_lead=None, style=None):
    p = doc.add_paragraph(style=style)
    rtl(p)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    rtl(p); p.paragraph_format.keep_with_next = True
    return p

def bullet(doc, text):
    p = para(doc, text, style="List Bullet")
    p.paragraph_format.right_indent = Cm(.65)
    p.paragraph_format.left_indent = Cm(0)
    return p

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); el = OxmlElement("w:shd")
    el.set(qn("w:fill"), fill); pr.append(el)

def borders(cell):
    pr = cell._tc.get_or_add_tcPr(); box = OxmlElement("w:tcBorders"); pr.append(box)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6"); el.set(qn("w:color"), "D9D9D9"); box.append(el)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,"244A3B"); borders(c)
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs: rtl(p)
        for r in c.paragraphs[0].runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    trpr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); trpr.append(repeat)
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); borders(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx%2: shade(cells[i],"F3F7F5")
            for p in cells[i].paragraphs: rtl(p)
            if widths: cells[i].width=Cm(widths[i])
    return t

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Cm(2); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2); sec.right_margin=Cm(2)

for name,size,before,after in [("Normal",11,0,6),("Title",25,0,14),("Heading 1",16,16,7),("Heading 2",13,11,5)]:
    s=doc.styles[name]; s.font.name="Arial"; s._element.rPr.rFonts.set(qn("w:cs"),"Arial")
    s._element.rPr.rFonts.set(qn("w:eastAsia"),"Arial"); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    if name != "Normal": s.font.bold=True

title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

title=doc.add_paragraph(style="Title"); rtl(title); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.add_run("תיק פרטיות ואבטחת מידע פנימי")
p=para(doc,"START Life Fit"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].bold=True; p.runs[0].font.size=Pt(15)
p=para(doc,"גרסה 1  עודכן ביום 24 בספטמבר 2026"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.color.rgb=RGBColor(90,90,90)
para(doc,"מסמך זה מרכז את נהלי העבודה הפנימיים של START Life Fit בנוגע לאיסוף, גישה, שמירה, מסירה ומחיקה של מידע אישי. מטרתו לאפשר ניהול עקבי ומתועד של מידע הלקוחות ולהפחית סיכוני פרטיות ואבטחה. המסמך מיועד לשימוש פנימי בלבד ואינו מיועד לפרסום באתר.")
para(doc,"בעל השליטה במידע: אלי כהן, עוסק מורשה 301658225. כתובת: לוי יצחק 10, דירה 2, בני ברק. דוא״ל: start.elicohenfitness@gmail.com. טלפון: 052-5290202.")

heading(doc,"1 אחריות וניהול המסמך",1)
para(doc,"אלי כהן אחראי לאישור הנהלים, לעדכונם ולתיעוד ביצועם. יש לעיין במסמך לפחות פעם בשנה וכן לאחר שינוי מהותי באפליקציה, בספק, בסוג המידע, בהרשאות או בעקבות אירוע אבטחה.")
table(doc,["פעולה","אחראי","תדירות","תיעוד"],[
    ["בדיקת רשימת בעלי הרשאה","אלי כהן","אחת לשישה חודשים","תאריך ותוצאה בטבלת המעקב"],
    ["בדיקת ספקים והרשאות","אלי כהן","אחת לשנה ובכל שינוי","רשימת ספקים מעודכנת"],
    ["בדיקת מחיקות","אלי כהן","אחת לרבעון","יומן בקשות ומחיקות"],
    ["בדיקת גיבוי ושחזור","אלי כהן או ספק טכני מורשה","אחת לשנה","סיכום הבדיקה"],
], [6,4,3,4])

heading(doc,"2 הגדרת המאגר ומטרותיו",1)
para(doc,"המאגר משמש לניהול שירותי ליווי כושר ותזונה, יצירת חשבונות, התאמת תוכניות, תיעוד אימונים וארוחות, מעקב התקדמות, תקשורת עם המאמן, תמיכה, אבטחה וחיוב. אין להשתמש במידע למטרה חדשה שאינה תואמת את ההודעה למשתמש ואת מדיניות הפרטיות בלי לבדוק תחילה את הצורך בעדכון ובהסכמה.")
table(doc,["קטגוריה","דוגמאות","רגישות","מטרה"],[
    ["זיהוי וקשר","שם דוא״ל טלפון ומזהה חשבון","רגיל","חשבון קשר ותמיכה"],
    ["בריאות וליווי","משקל מדידות מגבלות הערות וצעדים","רגישות מיוחדת","התאמת הליווי ומעקב"],
    ["תמונות וסרטונים","תמונות התקדמות ארוחות וסרטוני טכניקה","רגישות מיוחדת","משוב ומעקב"],
    ["אימון ותזונה","תוכניות ביצועים תפריטים ויומנים","אישי","ליווי והצגת התקדמות"],
    ["תקשורת","הודעות ותמיכה","אישי","מתן השירות וטיפול בבקשות"],
    ["טכני ושיווקי","יומנים טוקן התראות ונתוני Meta בהסכמה","משתנה","אבטחה תפעול ומדידה"],
], [3.5,6,3,4.5])

heading(doc,"3 מערכות וספקי עיבוד",1)
table(doc,["ספק","שימוש","מידע אפשרי","פעולת בקרה"],[
    ["Supabase","אימות מסד נתונים ואחסון","חשבון ליווי ומדיה","בדיקת הרשאות אחסון ומדיניות גישה"],
    ["Vercel","אירוח והפעלת שירותי ווב","יומני רשת ונתונים טכניים","הגבלת גישה לפרויקט וסודות"],
    ["OpenAI","הערכת ארוחה לפי בקשת המשתמש","תיאור או תמונת ארוחה","שליחה רק בהפעלת התכונה ובדיקת תנאי הספק"],
    ["Apple ו Google","בריאות והתראות","צעדים טוקן התראה ונתוני מכשיר","הרשאות מינימליות"],
    ["Open Food Facts","חיפוש ברקוד","ברקוד מוצר","אין לצרף מזהה משתמש"],
    ["YouTube","סרטוני הדרכה","נתוני צפייה טכניים","שימוש במצב מוגבר פרטיות ככל שניתן"],
    ["Meta","פיקסל בדף הנחיתה","צפייה והתחלת רכישה","הפעלה רק לאחר הסכמה"],
], [3.3,4.3,5,4.4])
para(doc,"לפני הוספת ספק חדש יש לתעד את מטרת השימוש, סוגי המידע, מיקום העיבוד, הרשאות, תקופת שמירה, אפשרות מחיקה ותנאי ההתקשרות. יש לשמור הסכם או תנאי עיבוד זמינים של הספק ולבחון אותם מחדש כאשר השירות משתנה.")

heading(doc,"4 ניהול הרשאות",1)
para(doc,"גישה למידע תינתן רק למי שזקוק לה לצורך תפקידו. לכל אדם יהיה חשבון אישי. אין לשתף סיסמאות או להשתמש בחשבון של אדם אחר. הרשאה תבוטל מיד בסיום תפקיד או כאשר אינה נדרשת עוד.")
table(doc,["שם","תפקיד","מערכת","רמת גישה","מועד אישור","מועד בדיקה או ביטול"],[
    ["אלי כהן","מפעיל ומאמן","כל המערכות","מנהל","________","________"],
    ["________","________","________","________","________","________"],
    ["________","________","________","________","________","________"],
], [3,3,3,3,3,3])
for x in ["להפעיל אימות רב שלבי בכל מערכת שתומכת בכך.","להשתמש בסיסמאות ייחודיות ובמנהל סיסמאות.","לבדוק הרשאות מנהל אחת לשישה חודשים.","לא להעביר קבצי לקוחות לחשבון אישי או למכשיר שאינו מוגן.","לנעול מכשירים בקוד ולהתקין עדכוני אבטחה."]: bullet(doc,x)

heading(doc,"5 שמירה ומחיקה",1)
para(doc,"כל עוד החשבון קיים, נתוני הליווי נשמרים כדי לאפשר חידוש עתידי בלי לאבד היסטוריה. כאשר מתקבלת בקשת מחיקה מאומתת, יש למחוק את חשבון האימות, הרשומות התלויות והמדיה הפרטית, למעט מידע שקיימת חובה חוקית לשמור או מידע הנחוץ זמנית להגנה על זכויות. יש לתעד את הבקשה, האימות, הפעולות ותאריך ההשלמה.")
for x in ["לקבל את הבקשה דרך התמיכה או הדוא״ל ולתעד את מועד קבלתה.","לאמת את זהות המבקש באופן מידתי בלי לאסוף מידע עודף.","לבדוק אם קיימת מניעה חוקית או מחלוקת פעילה המחייבת שמירה מוגבלת.","למחוק את המשתמש, נתוני הליווי והקבצים הפרטיים במערכות הפעילות.","לתעד אילו מערכות נבדקו, מי ביצע ומתי.","להודיע למבקש על השלמת הטיפול או על מידע שנשמר כדין."]: bullet(doc,x)
table(doc,["תאריך בקשה","משתמש","אימות","מערכות שנבדקו","תאריך השלמה","הערות"],[
    ["________","________","________","________","________","________"],
    ["________","________","________","________","________","________"],
], [3,3,3,4,3,3])

heading(doc,"6 בקשות עיון ותיקון",1)
para(doc,"בקשה לעיון או לתיקון תתקבל דרך התמיכה או הדוא״ל. יש לאמת את זהות המבקש, לאתר את המידע הרלוונטי, לבדוק זכויות של אחרים ולתעד את המענה. אין לשלוח מידע רגיש לערוץ שלא אומת.")
table(doc,["תאריך","סוג בקשה","מבקש","אופן אימות","החלטה ופעולה","מועד מענה"],[
    ["________","עיון תיקון מחיקה","________","________","________","________"],
    ["________","עיון תיקון מחיקה","________","________","________","________"],
], [3,3,3,4,5,3])

heading(doc,"7 קטינים והסכמת הורה",1)
para(doc,"חשבון של קטין יופעל רק לאחר קבלת הסכמה של הורה או אפוטרופוס. יש לתעד את שם ההורה, אמצעי קשר, הקשר לקטין, נוסח ההסכמה, התאריך ואופן קבלתה. אין לבקש מסמך מזהה אלא אם הוא נחוץ באופן סביר לאימות.")
table(doc,["שם הקטין","שם ההורה","קשר","טלפון או דוא״ל","תאריך הסכמה","אופן תיעוד"],[
    ["________","________","________","________","________","________"],
    ["________","________","________","________","________","________"],
], [3,3,2.5,4,3,3.5])

heading(doc,"8 אירוע אבטחת מידע",1)
para(doc,"אירוע אבטחה הוא חשיפה, אובדן, שינוי, מחיקה או גישה בלתי מורשית למידע, וכן חשד סביר לכך. יש לפעול מיד לצמצום הנזק ולא להסתיר או למחוק ראיות הנחוצות לבדיקה.")
for x in ["לעצור את הגישה או הפעולה המסוכנת בלי למחוק ראיות.","לתעד מועד גילוי, מערכת, סוג מידע ומספר משוער של אנשים שנפגעו.","להחליף סודות או הרשאות שנחשפו ולהפעיל ספק טכני לפי הצורך.","לבדוק אם נדרשת הודעה לרשות להגנת הפרטיות או לנושאי המידע.","לתעד החלטות, פעולות מתקנות ולקחים ולבדוק את יעילותן."]: bullet(doc,x)
table(doc,["תאריך ושעה","תיאור","מידע ומערכות","פעולת בלימה","דיווח והחלטה","סגירה"],[
    ["________","________","________","________","________","________"],
], [3,4,4,4,4,3])

heading(doc,"9 עבודה עם מידע ותמונות",1)
for x in ["לגשת רק למידע הדרוש למתן הליווי.","לא לשמור תמונות התקדמות או סרטונים בגלריה פרטית ללא צורך.","לא לשלוח מידע רפואי או תמונות גוף בקבוצות או לאדם שאינו מורשה.","לוודא את זהות הנמען לפני שליחת מידע אישי.","להימנע מהזנת מידע מזהה שאינו נחוץ לכלי בינה מלאכותית.","לא להשתמש בתמונת לקוח בפרסום או בעדות לקוח בלי הסכמה נפרדת ומתועדת."]: bullet(doc,x)

heading(doc,"10 עוגיות דיוור ופרסום",1)
para(doc,"Meta Pixel יופעל בדף הנחיתה רק לאחר הסכמה לעוגיות שאינן חיוניות. דחייה לא תמנע שימוש בדף. יש לאפשר שינוי בחירה ולשמור את הבחירה באחסון המקומי. הודעות שירות יופרדו מדיוור שיווקי. דיוור שיווקי יישלח רק כאשר קיים בסיס מתאים ותינתן דרך הסרה פשוטה.")
for x in ["לשמור תיעוד של הסכמה לדיוור, לרבות הנוסח והמועד.","לכבד בקשת הסרה בלי לשלוח פרסומות נוספות.","לראות בסיום התקשרות גם סירוב להמשך פרסום, ככל שנדרש לפי דין.","לא לכלול מידע בריאותי בקהלי פרסום או באירועי מדידה."]: bullet(doc,x)

heading(doc,"11 רשימת בדיקה תקופתית",1)
checks=["מדיניות הפרטיות תואמת את התכונות והספקים בפועל","Meta Pixel אינו נטען לפני הסכמה","הרשאות מנהל ובעלי גישה מעודכנות","אין קבצי לקוחות מיותרים במכשירים אישיים","בקשות מחיקה ועיון תועדו וטופלו","חשבונות של קטינים כוללים הסכמת הורה","ספקים חדשים נבדקו לפני שימוש","גיבויים והרשאות שחזור נבדקו","פרטי הקשר והקישורים באתר תקינים","אירועי אבטחה נבדקו והופקו לקחים"]
table(doc,["בדיקה","תקין","נדרש טיפול","תאריך והערה"],[[x,"☐","☐","________________"] for x in checks],[10,2,3,4])

heading(doc,"12 אישור ועדכון",1)
para(doc,"אני מאשר כי עיינתי במסמך, כי הוא משקף את אופן העבודה הידוע במועד החתימה, וכי אפעל לעדכנו כאשר יחול שינוי מהותי.")
table(doc,["שם","תפקיד","תאריך","חתימה"],[["אלי כהן","בעל העסק ובעל השליטה במידע","________","________"]],[5,5,4,4])

for section in doc.sections:
    h=section.header.paragraphs[0]; rtl(h); h.alignment=WD_ALIGN_PARAGRAPH.CENTER; h.text="START Life Fit  תיק פרטיות פנימי"
    for r in h.runs: r.font.name="Arial"; r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,100,100)
    f=section.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); f._p.append(field)

doc.core_properties.title="תיק פרטיות ואבטחת מידע פנימי START Life Fit"
doc.core_properties.subject="נהלי פרטיות אבטחת מידע והרשאות לשימוש פנימי"
doc.core_properties.author="START Life Fit"
doc.save(OUT)
print(OUT)
