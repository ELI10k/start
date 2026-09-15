from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path("outputs/START-legal-review-approval-form.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = alignment
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")
    for run in paragraph.runs:
        r_pr = run._r.get_or_add_rPr()
        run_rtl = r_pr.find(qn("w:rtl"))
        if run_rtl is None:
            run_rtl = OxmlElement("w:rtl")
            r_pr.append(run_rtl)
        run_rtl.set(qn("w:val"), "1")
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:cs"), "Arial")


def add_text(document, text, *, bold_lead=None, space_after=7, keep_with_next=False):
    paragraph = document.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.18
    paragraph.paragraph_format.keep_with_next = keep_with_next
    rtl(paragraph)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.keep_with_next = True
    rtl(paragraph)
    return paragraph


def add_check(document, text):
    paragraph = document.add_paragraph()
    paragraph.add_run("☐  ").bold = True
    paragraph.add_run(text)
    paragraph.paragraph_format.left_indent = Cm(0.25)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    rtl(paragraph)


def add_field_table(document, rows):
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].width = Inches(2.0)
        cells[1].width = Inches(4.7)
        cells[0].text = label
        cells[1].text = value
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                rtl(paragraph)
        cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(cells[0], "EAF4EE")
        if index % 2:
            set_cell_shading(cells[1], "F7F8F7")
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = document.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:cs"), "Arial")
styles["Normal"].font.size = Pt(11)
styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    if style_name == "Title":
        p_pr = style._element.get_or_add_pPr()
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)

title = document.add_paragraph("אישור בחינת מדיניות פרטיות ותנאי שימוש", style="Title")
title.paragraph_format.space_after = Pt(8)
title_pr = title._p.get_or_add_pPr()
title_border = title_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_pr.remove(title_border)
rtl(title)
subtitle = document.add_paragraph("אפליקציית START")
subtitle.runs[0].bold = True
subtitle.runs[0].font.size = Pt(14)
subtitle.paragraph_format.space_after = Pt(16)
rtl(subtitle)

add_text(
    document,
    "מסמך זה מיועד לתיעוד בדיקה משפטית של מדיניות הפרטיות ותנאי השימוש של אפליקציית START לפני פרסום מסחרי והגשה ל־App Store. המסמך אינו מהווה אישור משפטי כל עוד לא הושלמה הבדיקה והוא לא נחתם בידי עורך דין מוסמך.",
    space_after=12,
)

add_heading(document, "פרטי הבדיקה")
add_field_table(document, [
    ("שם השירות", "START"),
    ("מזהה אפליקציה", "co.il.startcoaching.app"),
    ("מפעיל השירות", "____________________________"),
    ("שם הישות המשפטית", "____________________________"),
    ("אימייל תמיכה", "start.elicohenfitness@gmail.com"),
    ("תאריך הבדיקה", "____ / ____ / ______"),
    ("גרסת המסמכים", "15 בספטמבר 2026"),
])

add_heading(document, "מסמכים שנבדקו")
add_check(document, "מדיניות הפרטיות בכתובת https://start.elicohenfitness.co.il/privacy")
add_check(document, "תנאי השימוש בכתובת https://start.elicohenfitness.co.il/terms")
add_check(document, "עמוד התמיכה בכתובת https://start.elicohenfitness.co.il/app-support")
add_check(document, "מיפוי הנתונים והצהרות הפרטיות המיועדות ל־App Store Connect")
add_check(document, "תהליך מחיקת החשבון והנתונים מתוך מסך הפרופיל")

add_heading(document, "היקף הבחינה")
add_text(document, "הבדיקה צריכה להתייחס לדין החל על מפעיל השירות ועל המשתמשים בפועל. הרשימה הבאה אינה מחליפה שיקול דעת משפטי והיא נועדה לוודא שנושאים מהותיים אינם נשמטים.")
for item in [
    "זהות מפעיל השירות, פרטי הקשר והגורם האחראי לטיפול בבקשות פרטיות.",
    "התאמת תיאור מטרות האיסוף והשימוש להתנהגות האפליקציה בפועל.",
    "הבסיס המתאים לאיסוף מידע אישי ורגיש וקבלת הסכמה במקום שבו היא נדרשת.",
    "איסוף משקל, מדידות, אלרגיות, הערות רפואיות, נתוני אימון ותזונה.",
    "קריאת צעדים מ־Apple Health והאיסור על שימוש בנתוני בריאות לפרסום או שיווק.",
    "איסוף תמונות התקדמות, תמונות מזון וסרטוני טכניקה והגישה של המאמן אליהם.",
    "העברת מידע לספקים ובהם Supabase, Vercel, OpenAI, Apple וספקי תוכן או התראות רלוונטיים.",
    "העברת מידע מחוץ לישראל, תנאי ההתקשרות עם הספקים ואמצעי ההגנה המתאימים.",
    "תקופות שמירה, גיבויים, יומני מערכת ותהליך מחיקת חשבון וקבצים.",
    "זכויות עיון, תיקון, מחיקה, ביטול הסכמה ופנייה בנושא פרטיות.",
    "אבטחת מידע, ניהול אירועי אבטחה והחובות הנובעות מרגישות המידע.",
    "התאמה לקטינים או קביעת גיל מינימלי ברור אם השירות אינו מיועד לקטינים.",
    "דיוק ההבהרה שהאפליקציה אינה מספקת אבחון או טיפול רפואי.",
    "התאמה לדיני הגנת הצרכן, להסכמי הליווי ולמדיניות תשלומים וביטולים, אם קיימים.",
    "התאמה להנחיות App Store ולתשובות שיוזנו בשאלון App Privacy.",
]:
    add_check(document, item)

add_heading(document, "תוצאת הבחינה")
add_text(document, "יש לסמן אפשרות אחת בלבד לאחר השלמת הבדיקה:")
add_check(document, "מאושר לפרסום בנוסח שנבדק, בכפוף לכך שהיישום והצהרות App Store Connect תואמים לנוסח.")
add_check(document, "מאושר בכפוף להשלמת התיקונים המפורטים במסמך זה ולאישור נוסח מעודכן.")
add_check(document, "אינו מאושר לפרסום בשלב זה.")

add_heading(document, "תיקונים או תנאים לאישור")
for number in range(1, 7):
    add_text(document, f"{number}. ____________________________________________________________________________", space_after=10)

document.add_page_break()
add_heading(document, "הצהרת עורך הדין")
add_text(
    document,
    "אני מאשר או מאשרת כי בחנתי את המסמכים והנושאים שסומנו לעיל על בסיס המידע שנמסר לי ועל בסיס הדין החל כפי שהבנתי אותו במועד הבדיקה. האישור מוגבל לגרסה ולנסיבות המפורטות במסמך זה. שינוי מהותי בסוגי המידע, במטרות השימוש, בספקים, בתשלומים, בקהל המשתמשים או בתכונות האפליקציה מחייב בחינה מחודשת.",
)
add_field_table(document, [
    ("שם עורך הדין", "____________________________"),
    ("מספר רישיון", "____________________________"),
    ("משרד", "____________________________"),
    ("אימייל", "____________________________"),
    ("חתימה", "____________________________"),
    ("תאריך", "____ / ____ / ______"),
])

add_heading(document, "אישור בעל השירות")
add_text(
    document,
    "אני מאשר או מאשרת כי מסרתי לבודק מידע מלא ומדויק על פעילות START, סוגי המידע, ספקי השירות, קהל המשתמשים, אמצעי התשלום ותהליך המחיקה. אני מתחייב או מתחייבת לעדכן את המסמכים ולבקש בחינה נוספת לפני שינוי מהותי.",
)
add_field_table(document, [
    ("שם מלא", "____________________________"),
    ("תפקיד", "____________________________"),
    ("חתימה", "____________________________"),
    ("תאריך", "____ / ____ / ______"),
])

footer = section.footer.paragraphs[0]
footer.text = "START  |  טופס בחינה משפטית  |  גרסת מסמכים 15 בספטמבר 2026"
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor(90, 95, 90)
rtl(footer, WD_ALIGN_PARAGRAPH.CENTER)

for paragraph in document.paragraphs:
    rtl(paragraph, paragraph.alignment or WD_ALIGN_PARAGRAPH.RIGHT)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT.resolve())
